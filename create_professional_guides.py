from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)

def border(cell, color='D9D9D9'):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:color'), color)

def set_cell_margin(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')

def set_run_font(run, name='Aptos', size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:ascii'), name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

def remove_paragraph_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn('w:pBdr'))
    if borders is not None:
        p_pr.remove(borders)

def style_document(doc, footer_text):
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Aptos'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in [('Title', 24), ('Heading 1', 16), ('Heading 2', 12.5), ('Heading 3', 11)]:
        s = styles[name]
        s.font.name = 'Aptos Display' if name == 'Title' else 'Aptos'
        s._element.rPr.rFonts.set(qn('w:ascii'), s.font.name)
        s._element.rPr.rFonts.set(qn('w:hAnsi'), s.font.name)
        s.font.size = Pt(size)
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.font.bold = name != 'Title'
        s.paragraph_format.space_before = Pt(16 if name == 'Heading 1' else 10)
        s.paragraph_format.space_after = Pt(6)
        p_pr = s._element.get_or_add_pPr()
        borders = p_pr.find(qn('w:pBdr'))
        if borders is not None:
            p_pr.remove(borders)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run(footer_text)
    set_run_font(r, size=8, color=(90, 90, 90))

def title(doc, main, subtitle):
    p = doc.add_paragraph(style='Title')
    remove_paragraph_border(p)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(main)
    set_run_font(r, name='Aptos Display', size=24, color=(0,0,0))
    p2 = doc.add_paragraph()
    r2 = p2.add_run(subtitle)
    set_run_font(r2, size=11, color=(70,70,70))
    p2.paragraph_format.space_after = Pt(16)

def add_paragraph(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p

def bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run_font(r, size=10.5)

def code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.12)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), 'F3F5F7')
    p_pr.append(shd)
    r = p.add_run(text)
    set_run_font(r, name='Courier New', size=8.5)
    return p

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ''
        r = hdr[i].paragraphs[0].add_run(h)
        set_run_font(r, size=9.5, bold=True, color=(255,255,255))
        shade(hdr[i], '1F4E78')
        border(hdr[i])
        set_cell_margin(hdr[i])
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row_index, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ''
            r = cells[i].paragraphs[0].add_run(str(value))
            set_run_font(r, size=9)
            if row_index % 2 == 1:
                shade(cells[i], 'F6F9FC')
            border(cells[i])
            set_cell_margin(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def h1(doc, text):
    p = doc.add_paragraph(text, style='Heading 1')
    return p

def h2(doc, text):
    return doc.add_paragraph(text, style='Heading 2')

def api_guide(path):
    doc = Document()
    style_document(doc, 'API Express PostgreSQL Professional Guide')
    title(doc, 'API Express PostgreSQL Professional Guide', 'A practical backend reference for building, testing, deploying, and explaining production Node.js APIs')
    add_paragraph(doc, 'This guide turns the project backend into a job-ready learning path. It explains the decisions behind an Express API, PostgreSQL and Neon connectivity, security boundaries, deployment variables, and the operational habits expected from a full-stack developer.', 'This guide')
    table(doc, ['Outcome', 'What you should be able to do'], [
        ['Design', 'Turn a UI requirement into stable resources, routes, request contracts, and response contracts.'],
        ['Build', 'Create modular Express routes, controllers, middleware, PostgreSQL queries, and real-time events.'],
        ['Operate', 'Configure Neon and Render safely with environment variables, health checks, CORS, logs, and repeatable schema setup.'],
        ['Interview', 'Explain trade-offs: REST vs RPC, JWT cookies vs headers, SQL constraints vs application checks, and pools vs single clients.'],
    ], [1.2, 5.2])

    h1(doc, '1 Backend architecture in this project')
    add_paragraph(doc, 'The server is an Express application running on Node.js. server.js creates the HTTP server, initializes Socket.IO, waits for table creation, and then listens on process.env.PORT. app.js configures middleware and mounts feature routes. Controllers implement request handling, while database/db.js owns the PostgreSQL pool.')
    table(doc, ['Layer', 'Responsibility', 'Project example'], [
        ['Bootstrap', 'Load configuration and start only after dependencies are ready.', 'server.js calls createTables before listen.'],
        ['Application', 'Set CORS, parsers, health route, and route mounting.', 'app.js registers /api/v1 routes.'],
        ['Route', 'Map HTTP method and path to a controller.', 'productRoutes exposes product actions.'],
        ['Controller', 'Validate input, call data layer, return an HTTP response.', 'productController filters, paginates, creates, and updates products.'],
        ['Database', 'Pool connections and issue parameterized SQL.', 'new Pool with DATABASE_URL or DB fields.'],
    ], [1.1, 2.6, 2.7])
    code(doc, "Client -> CORS -> Express middleware -> route -> controller -> PostgreSQL -> JSON response\n                         |                                  -> Socket.IO event when state changes")

    h1(doc, '2 Node project setup and runtime configuration')
    add_paragraph(doc, 'package.json describes scripts and dependencies. npm install installs runtime dependencies. npm run dev normally runs nodemon for local iteration; npm start runs the production entry point. Keep production dependencies in dependencies, not devDependencies.')
    code(doc, 'npm init -y\nnpm install express pg dotenv cors cookie-parser jsonwebtoken bcrypt\nnpm install --save-dev nodemon\nnpm run dev\nnpm start')
    add_paragraph(doc, 'Environment variables separate code from deployment-specific values. A local configuration file can set PORT=4000. Render supplies PORT automatically, so do not force 4000 in Render. Secrets never belong in client-side Vite variables or Git commits.')
    table(doc, ['Variable', 'Purpose', 'Where it belongs'], [
        ['DATABASE_URL', 'Neon PostgreSQL connection string.', 'Backend only: local config and Render environment.'],
        ['DB_SSL=true', 'Enables TLS for hosted PostgreSQL.', 'Backend only.'],
        ['JWT_SECRET_KEY', 'Signs authentication tokens.', 'Backend only, generated and rotated when exposed.'],
        ['FRONTEND_URL', 'Allow-list for browser origin in CORS.', 'Backend deployment environment.'],
        ['VITE_API_URL', 'Public API base URL used by React.', 'Frontend build environment only.'],
    ], [1.5, 2.4, 2.5])

    h1(doc, '3 HTTP APIs and REST contracts')
    add_paragraph(doc, 'An API is a contract between clients and the server. A REST-style API uses resource nouns in paths and HTTP methods to communicate intent. The route should be predictable, the response shape consistent, and error cases explicit.')
    table(doc, ['Method', 'Intent', 'Example', 'Expected status'], [
        ['GET', 'Read a collection or resource.', 'GET /api/v1/product?page=1', '200'],
        ['POST', 'Create a resource or action.', 'POST /api/v1/auth/register', '201'],
        ['PUT', 'Replace a complete resource.', 'PUT /api/v1/product/:id', '200'],
        ['PATCH', 'Partially update a resource.', 'PATCH /api/v1/order/:id/status', '200'],
        ['DELETE', 'Remove a resource.', 'DELETE /api/v1/product/:id', '200 or 204'],
    ], [0.7, 1.6, 2.7, 1.1])
    add_paragraph(doc, 'Use query parameters for optional filtering, sorting, and pagination. In this project, product search can accept availability, price, category, ratings, search, sort, and page. Validate allowed sort keys before building SQL; never concatenate unchecked user text into a query.')
    code(doc, "GET /api/v1/product?category=Beauty&price=0-10000&sort=price-low&page=1\n\n{\n  \"success\": true,\n  \"products\": [],\n  \"totalProducts\": 0,\n  \"currentPage\": 1,\n  \"totalPages\": 0\n}")

    h1(doc, '4 Express middleware and request lifecycle')
    add_paragraph(doc, 'Middleware runs in registration order. It can parse input, attach identity, reject an invalid request, or handle errors. A route handler should receive a request that has already passed the generic concerns it needs.')
    table(doc, ['Middleware', 'Why it exists', 'Common mistake'], [
        ['express.json()', 'Parses JSON bodies into req.body.', 'Registering it after JSON routes.'],
        ['cookieParser()', 'Reads signed or plain cookies.', 'Treating a cookie as proof of identity without verification.'],
        ['cors()', 'Controls which browser origins may call the API.', 'Using wildcard origin with credentials.'],
        ['auth middleware', 'Verifies JWT and attaches req.user.', 'Trusting a client role value.'],
        ['error middleware', 'Returns a consistent failure shape.', 'Forgetting next(error) from async handlers.'],
    ], [1.4, 2.6, 2.4])
    code(doc, "app.use(cors({ origin: allowedOrigin, credentials: true }));\napp.use(cookieParser());\napp.use(express.json({ limit: '1mb' }));\napp.use('/api/v1/product', productRoutes);\napp.use(errorMiddleware);")
    add_paragraph(doc, 'Stripe webhooks are a special case: signature verification requires the raw body. Register the webhook route with express.raw before express.json. This is a common production bug because JSON parsing changes the bytes Stripe signed.')

    h1(doc, '5 PostgreSQL foundations for backend developers')
    add_paragraph(doc, 'PostgreSQL is a relational database. Tables store rows; columns have types; constraints protect invariants. SQL is not only a query language. It is also a data integrity tool. Use it to prevent impossible states even if a client bypasses your UI.')
    table(doc, ['Concept', 'Meaning', 'Ecommerce example'], [
        ['Primary key', 'Unique identity for a row.', 'products.id as UUID.'],
        ['Foreign key', 'A row must reference an existing related row.', 'products.created_by references users.id.'],
        ['Unique constraint', 'Prevents duplicate values.', 'users.email and review product/user pair.'],
        ['Check constraint', 'Enforces a boolean rule.', 'price >= 0 and ratings between 0 and 5.'],
        ['Index', 'Speeds common lookup patterns.', 'Unique review index for product_id and user_id.'],
        ['Transaction', 'All statements succeed or none do.', 'Create an order, items, payment, and stock change together.'],
    ], [1.25, 2.55, 2.6])
    code(doc, "CREATE TABLE products (\n  id UUID PRIMARY KEY DEFAULT gen_random_uuid(),\n  name VARCHAR(255) NOT NULL,\n  price DECIMAL(10,2) NOT NULL CHECK (price >= 0),\n  category VARCHAR(100) NOT NULL,\n  stock INT NOT NULL CHECK (stock >= 0),\n  created_by UUID NOT NULL REFERENCES users(id)\n);")

    h1(doc, '6 Node pg pool and Neon connectivity')
    add_paragraph(doc, 'The pg Pool manages reusable PostgreSQL connections. A pool is safer than creating a connection per request and supports concurrent traffic. This project accepts a hosted DATABASE_URL for Neon while retaining an optional local DB_USER, DB_HOST, DB_NAME, DB_PASSWORD, and DB_PORT configuration path.')
    code(doc, "const ssl = process.env.DB_SSL === 'true' ? { rejectUnauthorized: false } : undefined;\nconst database = new Pool(\n  process.env.DATABASE_URL\n    ? { connectionString: process.env.DATABASE_URL, ssl }\n    : { user, host, database, password, port, ssl }\n);")
    add_paragraph(doc, 'Neon requires encrypted connections. Use DATABASE_URL and DB_SSL=true in the backend deployment. Do not split the Neon URL into localhost values. Never put it in Vercel because VITE variables are shipped to the browser.')
    h2(doc, 'Parameterized queries')
    add_paragraph(doc, 'Parameterized queries separate SQL structure from data. The database receives the text and the values independently, which prevents classic SQL injection. Parameterize values. For dynamic SQL fragments such as ORDER BY, use a fixed allow-list.')
    code(doc, "const result = await database.query(\n  'SELECT * FROM products WHERE category ILIKE $1 LIMIT $2 OFFSET $3',\n  [`%${category}%`, limit, offset]\n);")

    h1(doc, '7 Authentication authorization and session safety')
    add_paragraph(doc, 'Authentication answers who the user is. Authorization answers what that authenticated user may do. A safe API verifies the token on every protected request and checks role or ownership at the server, even when the frontend hides an admin button.')
    bullets(doc, [
        'Hash passwords with bcrypt before storing them. Never save plaintext passwords or send password hashes back to the client.',
        'Use a strong JWT secret and an expiration. Rotate it after exposure; an exposed secret invalidates trust in tokens signed with it.',
        'When using cookies across origins, set credentials: true in Axios and CORS, use exact HTTPS origins, and consider SameSite and Secure cookie flags.',
        'Authorize by resource ownership: a user may update only their own address, wishlist, or order unless their server-verified role is Admin.',
        'Return 401 for unauthenticated requests and 403 for authenticated but forbidden actions. Avoid revealing whether an email exists during password reset.',
    ])

    h1(doc, '8 Validation errors and observability')
    add_paragraph(doc, 'Validate at the boundary, then validate critical invariants again in the database. Client validation improves usability but is never a security boundary. A controller should reject malformed input early and return one consistent error shape.')
    table(doc, ['Failure type', 'Response', 'Example'], [
        ['Invalid input', '400 Bad Request', 'Negative price or missing required field.'],
        ['Unauthenticated', '401 Unauthorized', 'No valid access token.'],
        ['Forbidden', '403 Forbidden', 'Customer attempts admin operation.'],
        ['Missing resource', '404 Not Found', 'Unknown product ID.'],
        ['Conflict', '409 Conflict', 'Duplicate email or unique review.'],
        ['Unexpected failure', '500 Internal Server Error', 'Database or integration failure; log details server-side.'],
    ], [1.2, 1.8, 3.4])
    add_paragraph(doc, 'Add structured logs around request IDs, failed integrations, and startup checks. Do not log passwords, JWTs, database URLs, card data, or raw webhook secrets. A health route such as GET /api/v1/health should answer quickly and be safe to expose.')

    h1(doc, '9 CORS deployment and production checklist')
    add_paragraph(doc, 'CORS is enforced by browsers, not by curl. A backend can work from the terminal while the browser blocks it. The backend must allow the exact Vercel frontend origin and dashboard origin. Do not include a trailing slash if the configured origin check compares exact strings.')
    table(doc, ['Environment', 'API URL used by browser', 'Backend CORS value'], [
        ['Local', 'http://localhost:4000/api/v1', 'Allow local Vite origins during development.'],
        ['Vercel client', 'https://your-render-service.onrender.com/api/v1', 'FRONTEND_URL=https://your-vercel-site.vercel.app'],
        ['Dashboard', 'Same API base URL', 'DASHBOARD_URL=https://your-dashboard-site'],
    ], [1.3, 2.7, 2.4])
    bullets(doc, [
        'Render: Root Directory server, Build Command npm install, Start Command npm start.',
        'Render: add DATABASE_URL and DB_SSL=true; Render supplies PORT automatically.',
        'Vercel: add VITE_API_URL with /api/v1; rebuild after changing any VITE variable.',
        'Use HTTPS for deployed frontend origins and test a real browser request, not only curl.',
        'Keep .env and config files out of Git. Use environment dashboards or secret managers for production values.',
    ])

    h1(doc, '10 Testing troubleshooting and interview practice')
    h2(doc, 'Manual API test')
    code(doc, "curl 'http://localhost:4000/api/v1/product?price=0-10000&page=1' \\\n  -H 'Accept: application/json'")
    add_paragraph(doc, 'A 200 response with products: [] means the server works but the table has no matching records. A connection refused error usually means the server is not listening on that port. A browser-only CORS error means the API may be healthy but the origin allow-list is wrong.')
    h2(doc, 'Interview prompts')
    bullets(doc, [
        'Why use a PostgreSQL pool? Explain reuse, concurrency, resource limits, and graceful shutdown.',
        'How do you prevent SQL injection? Explain placeholders, allow-lists for SQL structure, and validation.',
        'Why are database constraints still necessary if the API validates input? Explain multiple clients and race conditions.',
        'How do cookies work across Vercel and Render? Explain credentials, exact CORS origins, HTTPS, SameSite, and Secure.',
        'How would you make checkout idempotent? Explain idempotency keys, transactions, webhook verification, and unique payment references.',
    ])
    h1(doc, '11 Portfolio standard for this backend')
    add_paragraph(doc, 'A hiring-ready project demonstrates more than endpoints. Document the API contract, include a seed or migration strategy, show protected routes, return clear errors, and deploy a working frontend. In a README, explain the architecture, local setup, environment variable names without values, test accounts only when safe, and the public health endpoint.')
    table(doc, ['Week', 'Practice milestone'], [
        ['1', 'Build a CRUD resource with parameterized queries, validation, pagination, and tests.'],
        ['2', 'Add registration, login, cookie/JWT verification, roles, and ownership checks.'],
        ['3', 'Build order creation with transactions, stock checks, and a fake payment provider.'],
        ['4', 'Deploy Express and Neon, configure CORS, add health checks, and write a clear architecture README.'],
    ], [0.8, 5.6])
    h2(doc, 'Production readiness check')
    bullets(doc, [
        'I can explain the API request flow from browser to controller to PostgreSQL and back.',
        'I can reproduce a local setup without committing secrets and can deploy it with environment variables.',
        'I can diagnose empty data, CORS failures, port conflicts, database connectivity errors, and invalid authentication.',
        'I can explain one security decision and one reliability decision in terms a non-specialist teammate understands.',
    ])
    doc.save(path)

def react_guide(path):
    doc = Document()
    style_document(doc, 'React Professional Full Stack Guide')
    title(doc, 'React Professional Full Stack Guide', 'A production focused React and Vite reference for full stack developers preparing for the job market')
    add_paragraph(doc, 'This guide explains React as a system for rendering state, composing reusable UI, managing server data, and delivering accessible production applications. It connects React concepts to the ecommerce client and dashboard in this project so each topic has a practical full-stack purpose.', 'This guide')
    table(doc, ['Career capability', 'Evidence in a strong project'], [
        ['UI architecture', 'Clear page, layout, feature, hook, and service boundaries.'],
        ['Data handling', 'Loading, empty, error, retry, caching, and mutation states are explicit.'],
        ['Security awareness', 'No secrets in VITE variables; UI role checks are backed by server authorization.'],
        ['Quality', 'Accessible controls, responsive layouts, tests for critical flows, and measurable performance.'],
    ], [1.7, 4.9])

    h1(doc, '1 React mental model')
    add_paragraph(doc, 'React renders a description of the interface from current props and state. When state changes, React calls components again and reconciles the result with the browser DOM. The main skill is not memorizing hooks; it is deciding what data exists, where it should live, and how the UI derives from it.')
    table(doc, ['Term', 'Professional interpretation', 'Ecommerce example'], [
        ['Component', 'A focused function that maps inputs to UI.', 'ProductCard renders one product.'],
        ['Props', 'Read-only inputs controlled by a parent.', 'ProductCard receives product and onAddToCart.'],
        ['State', 'Data that changes over time and causes rendering.', 'Search query, modal state, cart quantity.'],
        ['Derived data', 'Calculated from existing state; do not store it separately.', 'Cart total from items and quantities.'],
        ['Effect', 'Synchronizes React with an external system.', 'Fetch API data, subscribe to Socket.IO, sync local storage.'],
    ], [1.1, 3.1, 2.4])
    code(doc, "function CartSummary({ items }) {\n  const total = items.reduce((sum, item) => sum + item.price * item.quantity, 0);\n  return <p>AED {total.toFixed(2)}</p>;\n}")

    h1(doc, '2 JavaScript and TypeScript skills before advanced React')
    add_paragraph(doc, 'React code becomes clear when the JavaScript underneath is clear. Practice destructuring, array methods, modules, promises, async/await, object spread, optional chaining, and immutability. In job work, TypeScript adds compile-time contracts for props, API responses, and state.')
    table(doc, ['Skill', 'Why it matters in React'], [
        ['map filter find reduce', 'Transform server data into UI without mutation.'],
        ['Spread syntax', 'Create new arrays and objects for state updates.'],
        ['async await and try catch', 'Model network success, failure, and cleanup.'],
        ['Modules', 'Keep services, components, hooks, and utilities independently testable.'],
        ['Type narrowing', 'Safely handle nullable API values and union states.'],
    ], [2.0, 4.6])

    h1(doc, '3 Vite project structure and environment variables')
    add_paragraph(doc, 'Vite provides a fast development server and production build. main.jsx mounts the React tree. App.jsx usually owns top-level providers and route layout. Keep feature code close to the feature, while cross-cutting infrastructure lives in lib, store, hooks, and shared components.')
    code(doc, "src/\n  pages/          // route-level screens\n  components/     // reusable visual units\n  layouts/        // navigation and page shells\n  lib/            // axios, realtime, utilities\n  store/          // Redux store and slices\n  contexts/       // lightweight shared UI state\n  hooks/          // reusable behavior\n  data/           // safe fallback display data")
    add_paragraph(doc, 'Vite only exposes variables prefixed with VITE_. That makes them public at build time. VITE_API_URL is correct for a public API base URL. DATABASE_URL, JWT_SECRET_KEY, Stripe secret keys, SMTP passwords, and Cloudinary secrets must never be included in a Vite environment file.')
    table(doc, ['Variable', 'Correct use'], [
        ['VITE_API_URL', 'https://your-render-service.onrender.com/api/v1 in Vercel; localhost API during local work.'],
        ['VITE_STRIPE_PUBLISHABLE_KEY', 'Browser-safe Stripe publishable key only.'],
        ['VITE_GOOGLE_MAPS_API_KEY', 'Browser-restricted Google Maps key with allowed origins.'],
        ['VITE_DASHBOARD_URL', 'Public dashboard link for admin navigation.'],
    ], [2.1, 4.5])

    h1(doc, '4 JSX components props and composition')
    add_paragraph(doc, 'JSX is JavaScript syntax for describing elements. Use className, pass expressions in braces, and keep a component responsible for one coherent UI concern. Composition is the main scaling tool: a ProductPage can compose gallery, price, reviews, cart action, and recommendation sections without one giant component.')
    code(doc, "function ProductCard({ product, onAdd }) {\n  return (\n    <article>\n      <img src={product.images?.[0]?.url} alt={product.name} />\n      <h2>{product.name}</h2>\n      <button onClick={() => onAdd(product)}>Add to cart</button>\n    </article>\n  );\n}")
    bullets(doc, [
        'Use a stable database ID as a list key. Array indexes cause bugs when rows are reordered, filtered, or deleted.',
        'Do not mutate props. A child requests a change through a callback; the owner of the state performs the update.',
        'Use children for layout primitives such as Modal, Dialog, PageShell, and Card rather than creating many nearly identical wrappers.',
        'Avoid boolean prop explosion. When variants become complex, consider explicit composition or a small variant API.',
    ])

    h1(doc, '5 State events forms and reducers')
    add_paragraph(doc, 'State should be minimal and owned by the lowest common parent that needs it. Store source-of-truth values, not values that can be calculated from them. Use functional state updates when the next value depends on the previous value.')
    code(doc, "setCart((current) =>\n  current.some((item) => item.id === product.id)\n    ? current.map((item) => item.id === product.id ? { ...item, quantity: item.quantity + 1 } : item)\n    : [...current, { ...product, quantity: 1 }]\n);")
    table(doc, ['Approach', 'Use it when', 'Avoid it when'], [
        ['useState', 'A screen has a few independent values.', 'Many transitions must stay coordinated.'],
        ['useReducer', 'Actions describe related state changes.', 'State is simple enough for one setter.'],
        ['Context', 'A small shared concern crosses many levels.', 'High-frequency large data causes broad rerenders.'],
        ['Redux Toolkit', 'Complex app-wide client state needs predictable slices.', 'You only need local screen state.'],
    ], [1.4, 2.8, 2.3])
    h2(doc, 'Forms')
    add_paragraph(doc, 'Controlled inputs keep display value and React state aligned. Use semantic labels, appropriate input types, clear errors, disabled submit state while saving, and server validation. For complex forms, React Hook Form plus Zod can reduce boilerplate while keeping validation explicit.')

    h1(doc, '6 Effects server data and Axios')
    add_paragraph(doc, 'useEffect is for synchronization with systems outside React: network requests, subscriptions, timers, browser APIs, and imperative libraries. Do not use an effect to calculate a total from existing state; calculate it in render or memoize only when profiling proves a cost.')
    code(doc, "const axiosInstance = axios.create({\n  baseURL: import.meta.env.VITE_API_URL || 'http://localhost:4000/api/v1',\n  withCredentials: true,\n});\n\nconst { data } = await axiosInstance.get('/product', { params: { page: 1, sort: 'newest' } });")
    table(doc, ['State', 'UI behavior'], [
        ['loading', 'Show a skeleton or progress state without hiding page structure.'],
        ['success with data', 'Render the list, summary, and controls.'],
        ['success empty', 'Explain that no products match and offer a clear filter reset.'],
        ['error', 'Show a retry action and log useful diagnostic context.'],
        ['saving', 'Prevent duplicate mutation submissions; preserve user input on failure.'],
    ], [1.4, 5.1])
    add_paragraph(doc, 'Cancel or ignore obsolete requests when a user changes search terms quickly or navigates away. A custom hook such as useProducts can own the request, status, retry action, and cancellation so page components remain focused on UI.')

    h1(doc, '7 Routing navigation and protected screens')
    add_paragraph(doc, 'React Router maps locations to elements. Route components are screen-level units. Use nested layouts for shared navigation, URL parameters for resource identity, and search parameters for filters that should be shareable and bookmarkable.')
    code(doc, "<Routes>\n  <Route element={<StoreLayout />}>\n    <Route path=\"/\" element={<Home />} />\n    <Route path=\"/products\" element={<Products />} />\n    <Route path=\"/products/:productId\" element={<ProductDetail />} />\n  </Route>\n  <Route path=\"/orders\" element={<RequireAuth><Orders /></RequireAuth>} />\n</Routes>")
    add_paragraph(doc, 'A frontend protected route improves navigation but does not secure data. The server must verify the cookie or token and role on each protected API request. For example, hiding the admin dashboard link is only convenience; backend authorization is the real control.')

    h1(doc, '8 Redux client state and real time updates')
    add_paragraph(doc, 'This project uses a Redux store with slices for auth, cart, products, wishlist, popup, storefront, and admin state. A slice should contain actions and reducers for one domain. Keep normalized state where possible and avoid duplicating server data in many unrelated slices.')
    code(doc, "const cartSlice = createSlice({\n  name: 'cart',\n  initialState: { items: [] },\n  reducers: {\n    itemAdded(state, action) { state.items.push(action.payload); }\n  }\n});")
    add_paragraph(doc, 'For remote data that needs caching, retries, invalidation, or optimistic updates, use a server-state cache library such as TanStack Query. Socket.IO is useful when the UI needs to react to server events such as product changes, order updates, or admin changes. Subscribe in an effect, clean up listeners on unmount, and still fetch the authoritative server state after important mutations. An event is a notification, not a replacement for authorization or durable data.')

    h1(doc, '9 Performance and rendering discipline')
    table(doc, ['Technique', 'Use when', 'Warning'], [
        ['React.memo', 'A pure child rerenders often with unchanged props.', 'Measure first; unstable object props defeat it.'],
        ['useMemo', 'A costly derived calculation is proven expensive.', 'It is not a default performance tool.'],
        ['useCallback', 'A stable callback matters to a memoized child or effect.', 'It can add noise without benefit.'],
        ['Code splitting', 'A route or heavy feature should load on demand.', 'Keep loading fallback usable.'],
        ['Image optimization', 'Product images dominate page weight.', 'Use descriptive alt text and avoid layout shifts.'],
        ['List virtualization', 'Hundreds or thousands of rows render at once.', 'Do not virtualize small product grids prematurely.'],
    ], [1.5, 2.7, 2.3])
    add_paragraph(doc, 'Use browser DevTools and React DevTools profiler to identify actual bottlenecks. Optimize perceived performance too: fast initial route rendering, skeletons, stable image dimensions, responsive controls, and useful empty/error messages.')

    h1(doc, '10 Accessibility security and production quality')
    bullets(doc, [
        'Use semantic landmarks, headings in order, buttons for actions, links for navigation, and labels for every form control.',
        'Support keyboard navigation, visible focus states, escape-to-close dialogs, and meaningful image alt text.',
        'Never render untrusted HTML without sanitization. React escapes text by default; dangerous HTML bypasses that protection.',
        'Treat all VITE variables as public. Keep database, JWT, email, and payment secrets on the backend only.',
        'Use error boundaries for rendering failures and a route-level fallback that helps users recover.',
        'Test responsive layouts on narrow touch screens, not only a desktop browser window.',
    ])
    h2(doc, 'CORS and authentication integration')
    add_paragraph(doc, 'When React calls a different origin, Axios needs withCredentials: true for cookie authentication. The backend must send Access-Control-Allow-Credentials and allow the exact frontend origin. If a terminal request works but the browser says Browser origin is not allowed, inspect Render FRONTEND_URL and Vercel VITE_API_URL.')

    h1(doc, '11 Testing and debugging workflow')
    table(doc, ['Level', 'What to test', 'Example'], [
        ['Unit', 'Pure functions and reducers.', 'Cart quantity, price total, filter builder.'],
        ['Component', 'Interaction and accessibility behavior.', 'Login form error, disabled submit, modal focus.'],
        ['Integration', 'Screen plus mocked API.', 'Products page loading then empty state.'],
        ['End to end', 'Critical browser journey.', 'Register, login, add product, checkout, view order.'],
    ], [1.0, 2.9, 2.6])
    add_paragraph(doc, 'Debug from the boundary inward: verify the browser request URL and status, inspect the Network response body, confirm VITE_API_URL at build time, check backend logs, then query the database only when appropriate. Do not assume an empty screen means React is broken; it may be a valid empty API result.')

    h1(doc, '12 Job market roadmap and interview stories')
    add_paragraph(doc, 'For interviews, be able to tell a concrete story about your project: the user problem, component boundaries, API contract, authentication model, deployment issue you solved, and a measurable quality improvement. Recruiters value clear reasoning over a long list of libraries.')
    table(doc, ['Portfolio milestone', 'What to demonstrate'], [
        ['Product catalogue', 'Filtering, pagination, loading/empty/error states, responsive cards, accessible images.'],
        ['Authentication', 'Login/logout, persisted session, protected routes, backend authorization explanation.'],
        ['Cart and checkout', 'Reducer or slice, optimistic UI boundaries, price calculation, server-side verification.'],
        ['Admin dashboard', 'Role-aware navigation, CRUD forms, validation, real-time refresh or refetch strategy.'],
        ['Deployment', 'Vercel environment configuration, Render API URL, Neon database, CORS troubleshooting.'],
    ], [1.8, 4.7])
    h2(doc, 'Interview prompts')
    bullets(doc, [
        'Explain why useEffect is not the right place for derived data.',
        'Compare Context, Redux Toolkit, and a server-state cache library.',
        'Describe how you would prevent a stale search response from replacing a newer one.',
        'Explain why a VITE database URL would be a security incident.',
        'Describe how frontend role checks and backend authorization work together.',
        'Show how you investigate a 200 response with an empty products array.',
    ])
    h2(doc, 'Weekly practice routine')
    bullets(doc, [
        'Build one screen from a written user story, including loading, empty, error, and mobile states.',
        'Refactor one repeated UI pattern into a tested component with a small, clear prop contract.',
        'Trace one browser request through Vite configuration, Axios, Express, and the database response.',
        'Practice a two-minute explanation of one project decision: state ownership, CORS, authentication, or performance.',
    ])
    doc.save(path)

api_guide(ROOT / 'API_Express_PostgreSQL_Professional_Guide.docx')
react_guide(ROOT / 'React_JS_Professional_FullStack_Guide.docx')
